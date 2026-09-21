import contextlib
import io
import os
from typing import Dict, List

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from data_io.public_sheets_reader import (
    load_participants_from_public_form_sheet,
    load_participants_from_public_sheet,
)
from data_io.sheets_manager import (
    load_participants_from_form_sheet,
    load_participants_from_sheet,
    save_plan_to_sheet,
)
from logic.milp_allocator_v3 import generate_full_plan_cpsat, DEFAULT_TIME_LIMIT
from logic.back_sim import BackSimConfig, run_back_sim, summarize_conditions, check_priority_sections
from logic.car_pool import section_label, LARGE_CAR_IDS, NORMAL_CAR_IDS
from data_io.output_writer import (
    write_plan_xlsx,
    write_repaired_xlsx,
    write_updated_output_xlsx,
    compute_individual_summary_static,
    describe_car_row_changes,
    SECTION_COLORS,
)
from data_io.plan_importer import (
    fetch_output_format_xlsx_from_google_sheet,
    import_participants_from_output_xlsx,
    import_plan_from_output_xlsx,
)
from validator import (
    validate_participants,
    validate_transitions,
    count_car_changes,
    compute_runner_satisfaction,
    compute_individual_summary,
)

CREDENTIALS_PATH = "credentials.json"
OUTPUT_XLSX_PATH = "hakone_result.xlsx"
BACK_OUTPUT_XLSX_PATH = "hakone_result_back.xlsx"
REPAIR_OUTPUT_XLSX_PATH = "hakone_result_repaired.xlsx"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

APP_VERSION = "v3-cpsat-2026-09-17"

# 裏シュミを開くための隠しドット。ページ上の複数箇所に見た目が同じドットを散らし、
# この順番(dot_2 → dot_4 → dot_1)でクリックしたときだけ開閉する。それ以外のドット
# (dot_3など)や順番を間違えたクリックは進捗をリセットする「罠」として機能する。
_HIDDEN_SEQUENCE = ["dot_2", "dot_4", "dot_1"]


def _register_hidden_click(spot_id: str) -> None:
    progress = st.session_state.get("hidden_click_progress", 0)
    expected = _HIDDEN_SEQUENCE[progress] if progress < len(_HIDDEN_SEQUENCE) else None
    if spot_id == expected:
        progress += 1
        if progress >= len(_HIDDEN_SEQUENCE):
            st.session_state.show_back_sim = not st.session_state.get("show_back_sim", False)
            progress = 0
        st.session_state.hidden_click_progress = progress
    else:
        # 間違えても、その一手がたまたま最初の一手と一致するならそこからやり直せる
        # ようにする(そうしないと「1手目」を忘れて先頭から全部押し直す羽目になる)。
        st.session_state.hidden_click_progress = 1 if spot_id == _HIDDEN_SEQUENCE[0] else 0


def _hidden_dot(spot_id: str, key: str) -> None:
    """裏シュミを開くための隠しドットを1つ描画する。見た目は他のドットと区別が
    つかない、右端に寄せた小さい灰色の「・」。"""
    st.markdown(
        f"""
        <style>
        div.st-key-{key} button {{
            background: transparent; border: none; color: #cccccc;
            font-size: 11px; padding: 0 2px; min-height: 0; line-height: 1.2;
        }}
        div.st-key-{key} button:hover {{ color: #999999; background: transparent; border: none; }}
        div.st-key-{key} button:focus:not(:active) {{ color: #999999; box-shadow: none; }}
        </style>
        """,
        unsafe_allow_html=True,
    )
    _spacer, _dot_col = st.columns([40, 1])
    with _dot_col:
        if st.button("·", key=key):
            _register_hidden_click(spot_id)


def _car_id_of_cell(text):
    """個人別まとめの1セルのテキスト(例: "🚘 運転手(L1)")から車両IDを取り出す。
    車に乗っていない役割(ランナー・現地集合・不参加)ならNone。「運転手」「同乗者」の
    直後の括弧だけを見るので、修理の変更箇所ハイライトで末尾に「（旧: ...）」が
    追記されていても正しく車IDだけを取り出せる。⚠️重複(運転手と同乗者を兼ねる等の
    異常データ)では複数マッチしうるが、元のテキスト生成順(compute_individual_summary
    ではランナー→運転手→同乗者の順)の最後、すなわち一番優先度の低い役割の車IDを
    採用する(旧: ...を追記する前の、文字列末尾を見ていた挙動に合わせるため)。"""
    import re
    if not isinstance(text, str):
        return None
    matches = re.findall(r"(?:運転手|同乗者)\(([^)]+)\)", text)
    return matches[-1] if matches else None


def _condition_cell_styles(table: pd.DataFrame, together_map: dict, apart_map: dict, labels: list) -> pd.DataFrame:
    """一緒にしたい/離したい条件が付いた人について、名前欄だけでなく、実際にその
    条件を満たせている区間のセルにも同じ色・太字を適用するためのスタイル表を作る
    (Styler.apply(axis=None)用。tableは元データそのまま、返り値は同じ形のCSS文字列表)。
    「一緒にしたい」は同じ車になれていれば満たしている、「離したい」は指定した
    相手全員と同じ車になっていなければ満たしている、とみなす。"""
    name_to_row = {name: idx for idx, name in zip(table.index, table["名前"])}
    styles = pd.DataFrame("", index=table.index, columns=table.columns)

    for idx, row in table.iterrows():
        name = row["名前"]
        together_targets = together_map.get(name, set())
        apart_targets = apart_map.get(name, set())
        if not together_targets and not apart_targets:
            continue
        for label in labels:
            my_car = _car_id_of_cell(row[label])
            if my_car is None:
                continue

            def _shares_car_with(target_name):
                t_idx = name_to_row.get(target_name)
                if t_idx is None:
                    return False
                return _car_id_of_cell(table.at[t_idx, label]) == my_car

            if together_targets and any(_shares_car_with(t) for t in together_targets):
                styles.at[idx, label] = "color: #CC3333; font-weight: bold"
            elif apart_targets and not any(_shares_car_with(t) for t in apart_targets):
                styles.at[idx, label] = "color: #1E5FCC; font-weight: bold"

    return styles


def _render_individual_summary_table(
    plan, participants, key_prefix, highlight=None, summary=None, old_summary=None,
    unapplied_keys=None,
):
    """個人別まとめ(区間ごとの役割)の表を描画する。_render_plan_result(表シュミ・裏シュミの
    結果表示)と修理の結果表示の両方から使う共通部品。summaryを渡さなければ
    (表シュミ・裏シュミと同じ)validator.compute_individual_summaryで計算する。
    修理の結果表示では、実際にExcelに書き込まれる内容(重複検知込み)と一致させるため、
    output_writer.compute_individual_summary_staticの結果を渡す(修理で「適用しない」に
    された分もすべて適用したフルの内容。実際に出力ファイルへ反映するかはunapplied_keys
    で見た目だけ変える)。
    old_summary(名前をキーにした修理前の表示値。_load_old_individual_summary_by_name
    の結果)を渡すと、修理で変わったセルに(旧: ...)を追記する(フルスクリーン表示でも
    見えるように、ツールチップではなくセルのテキスト自体に埋め込む)。
    unapplied_keys((名前,区間)のset)に含まれるセルは、修理で変わってはいるが
    「適用しない」と選ばれているため、(旧: ...)の表示はそのままに、赤字太字の
    強調だけ外す(実際の出力ファイルではこのセルは修理前の値のまま書き出される)。"""
    if summary is None:
        summary = compute_individual_summary(plan, participants)
    unapplied_keys = unapplied_keys or set()
    labels = [section_label(s) for s in list(range(1, 11)) + [11]]
    table = pd.DataFrame(
        [
            {"名前": p.name, **{label: summary.get(pid, {}).get(label, "") for label in labels}}
            for pid, p in participants.items()
        ]
    )

    changed_mask = pd.DataFrame(False, index=table.index, columns=table.columns)
    if old_summary is not None:
        for idx, name in table["名前"].items():
            old_row = old_summary.get(name, {})
            for label in labels:
                old_val = old_row.get(label) or ""
                new_val = table.at[idx, label] or ""
                if old_val != new_val:
                    changed_mask.at[idx, label] = True
                    table.at[idx, label] = f"{new_val}　（旧: {old_val or '空欄'}）"

    def _cell_color(text):
        if isinstance(text, str) and text.startswith("⚠️重複"):
            return "background-color: #FF9999"
        if isinstance(text, str) and text.startswith("現地集合"):
            return "background-color: #D9B3FF"
        if isinstance(text, str) and text.startswith("🏃"):
            return "background-color: #FFE699"
        if isinstance(text, str) and text.startswith("🚘"):
            return "background-color: #BDD7EE"
        if isinstance(text, str) and text.startswith("👥"):
            return "background-color: #C6E0B4"
        return "background-color: #FFFFFF"

    styled_table = table.style.map(_cell_color, subset=labels)
    caption = "🟡 ランナー　🔵 運転手　🟢 同乗者　🟣 現地集合(1区は走らず2区から参加)　⚪ 不参加"

    if changed_mask.to_numpy().any():
        applied_mask = pd.DataFrame(True, index=table.index, columns=table.columns)
        for idx, name in table["名前"].items():
            for label in labels:
                if (name, label) in unapplied_keys:
                    applied_mask.at[idx, label] = False
        border_styles = pd.DataFrame("", index=table.index, columns=table.columns)
        for idx in table.index:
            for label in labels:
                if changed_mask.at[idx, label] and applied_mask.at[idx, label]:
                    border_styles.at[idx, label] = "color: #CC3300; font-weight: bold"
        styled_table = styled_table.apply(lambda _: border_styles, axis=None)
        caption += (
            "　｜　✏️ 赤文字＝修理で内容が変わり適用されるセル"
            "　｜　（旧:...）だけ表示＝変わったが適用しないセル（内容は修理前の値のまま）"
        )

    if highlight:
        together_names = highlight.get("together_names", set())
        apart_names = highlight.get("apart_names", set())
        driver_range_names = highlight.get("driver_range_names", set())

        def _name_color(name):
            in_together = name in together_names
            in_apart = name in apart_names
            if in_together and in_apart:
                return "color: #A000A0; font-weight: bold"  # 両方指定されている人は紫で目立たせる
            if in_apart:
                return "color: #1E5FCC; font-weight: bold"
            if in_together:
                return "color: #CC3333; font-weight: bold"
            return ""

        def _bold_driver_row(row):
            if row["名前"] not in driver_range_names:
                return ["" for _ in row.index]
            return [
                "font-weight: bold" if (col in labels and isinstance(row[col], str) and row[col].startswith("🚘")) else ""
                for col in row.index
            ]

        styled_table = styled_table.map(_name_color, subset=["名前"])
        styled_table = styled_table.apply(_bold_driver_row, axis=1)

        together_map = highlight.get("together_map", {})
        apart_map = highlight.get("apart_map", {})
        if together_map or apart_map:
            styled_table = styled_table.apply(
                _condition_cell_styles, axis=None,
                together_map=together_map, apart_map=apart_map, labels=labels,
            )

        caption += (
            "　｜　🔴 一緒にしたい人／一緒になれた区間　🔵 離したい人／離せた区間"
            "　🟣 両方指定　太字の🚘=運転区間数を指定した人"
        )

    st.dataframe(styled_table, hide_index=True, use_container_width=True, key=f"{key_prefix}_table")
    st.caption(caption)


def _render_plan_result(plan, participants, xlsx_path, key_prefix, download_label, highlight=None):
    """表シュミ・裏シュミの結果表示で共通する部分(区間ごとの配車・個人別まとめ・
    メトリクス・Excelダウンロード)をまとめた描画関数。
    highlight: 裏シュミの条件を個人別まとめ表で見分けやすくするための追加情報
    (dict: together_names/apart_names/driver_range_names のset、
     together_map/apart_map の 人名->相手名setで区間セルごとの充足判定に使う)。
    Noneなら何も強調しない。"""
    used_cars = len({car.car_id for section in plan for car in section.cars})
    large_n = len({car.car_id for section in plan for car in section.cars if car.car_type == "large"})
    normal_n = used_cars - large_n
    st.success(f"計算完了！　大型 {large_n} 台 ＋ 普通 {normal_n} 台 ＝ 合計 {used_cars} 台")

    for section in plan:
        label = section_label(section.section_id)
        runners = [participants[pid].name for pid in section.runner_ids if pid in participants]
        runner_text = f"🏃 ランナー {len(runners)} 名: {', '.join(runners)}" if runners else "（ランナーなし）"

        with st.expander(f"【{label}】　{runner_text}"):
            for car in section.cars:
                driver_name = participants[car.driver_id].name if car.driver_id in participants else "エラー"
                passengers = [participants[pid].name for pid in car.passenger_ids if pid in participants]
                car_type = "大型" if car.car_type == "large" else "普通"
                mt_badge = "　⛰️ 山行き" if car.is_mountain_goer else ""
                hotel_badge = "　🏨 ホテル組" if car.group == "hotel" else ""
                adv_badge = "　🚀 先行" if car.is_advance else ""

                st.markdown(f"**🚘 車 {car.car_id}**（{car_type}）{mt_badge}{hotel_badge}{adv_badge}")
                st.write(f"　👨‍✈️ 運転手: {driver_name}")
                st.write(f"　👥 同乗者: {', '.join(passengers) if passengers else 'なし'}")
                st.divider()

    with st.expander("👤 個人別まとめ（区間ごとの役割）"):
        _render_individual_summary_table(plan, participants, key_prefix, highlight=highlight)

    with st.expander("🚘 区間別配車（Excelの「区間別配車」シートと同じ形式）"):
        _render_car_assignment_table(plan, participants, key_prefix=key_prefix)

    transition_errors = validate_transitions(plan, participants)
    if transition_errors:
        st.warning("⚠️ 車両引き継ぎに問題があります: " + " / ".join(transition_errors))

    changes = count_car_changes(plan)
    n_wanted, n_satisfied, total_wanted, total_ran = compute_runner_satisfaction(plan, participants)
    col1, col2 = st.columns(2)
    col1.metric("ランナー希望充足", f"{n_satisfied}/{n_wanted}人", help="希望区間数を達成できた人の割合")
    col2.metric("走行区間数", f"{total_ran}/{total_wanted}区間", help="実際に走れた区間 / 希望合計")
    st.caption(f"ℹ️ 区間をまたいで車を乗り換えた回数: {changes}回")

    if os.path.exists(xlsx_path):
        with open(xlsx_path, "rb") as f:
            st.download_button(
                download_label,
                f,
                file_name=os.path.basename(xlsx_path),
                mime=XLSX_MIME,
                type="primary",
                key=f"{key_prefix}_download",
            )


# 区間ごとの色分け(1〜10区)。ユーザー指定の配色見本から実際のピクセル色を抽出した値。
def _render_section_color_legend():
    """「次に走る区間」の色分け(SECTION_COLORS)が、どの区間がどの色かを示す凡例。"""
    chips = "".join(
        f'<span style="background-color:{SECTION_COLORS[s]}; padding:2px 10px; '
        f'margin-right:6px; border-radius:4px; display:inline-block; font-size:0.85em;">{s}区</span>'
        for s in range(1, 11)
    )
    st.markdown(
        f'<div>🎨 セルの色（運転手・同乗者が次に走る区間）: {chips}</div>',
        unsafe_allow_html=True,
    )


def _render_car_assignment_table(plan, participants, key_prefix, old_table=None, revert_car_keys=None):
    """Excelの「区間別配車」シートと同じ形式(1区間=1ブロック、車ID・車種・山行き・
    先行・運転手・同乗者…)の表を画面上にも出す。列構成・基本の表示テキストは
    output_writer側の_car_blocks/_cars_headerをそのまま再利用し、「先行」列の上書き・
    セルの背景色もoutput_writer.compute_car_table_overridesを共用することで、
    Excel出力(区間別配車シート)と画面表示が完全に一致するようにしている。
    old_table((区間ラベル,車ID)をキーにした修理前の値。_load_old_cars_tableの結果)を
    渡すと、修理で変わったセルに(旧: ...)を追記する。
    revert_car_keys((区間ラベル,車ID)のset)に含まれる車は、UIで「適用しない」と
    選ばれ実際の出力ファイルでは修理前の値のまま書き出されるが、この表示は
    (旧: ...)の表示はそのままに赤字太字の強調だけ外す(出力ファイルの内容と
    見分けられるように、内容自体は変えない)。"""
    from data_io.output_writer import _car_blocks, _cars_header, compute_car_table_overrides

    blocks = _car_blocks(plan, participants)
    header = _cars_header(blocks)
    n_cols = len(header)
    advance_col = header.index("先行")
    revert_car_keys = revert_car_keys or set()

    rows = []
    row_keys = []  # (区間ラベル, 車ID) を各行に対応させたリスト(車が無い区間はcar_id=None)
    for label, section_rows in blocks:
        if not section_rows:
            rows.append([label] + [""] * (n_cols - 1))
            row_keys.append((label, None))
            continue
        for i, row in enumerate(section_rows):
            padded = list(row) + [""] * (n_cols - 1 - len(row))
            rows.append([label if i == 0 else ""] + padded)
            row_keys.append((label, row[0]))

    advance_overrides, cell_fills, pickup_cols = compute_car_table_overrides(plan, participants, header)
    for i, override in enumerate(advance_overrides):
        if override:
            rows[i][advance_col] = override

    table = pd.DataFrame(rows, columns=header).fillna("")

    changed_mask = pd.DataFrame(False, index=table.index, columns=table.columns)
    if old_table is not None:
        for i, (label, car_id) in enumerate(row_keys):
            if car_id is None:
                continue
            old_row = old_table.get((label, car_id))
            if old_row is None:
                continue
            for col_name in header:
                if col_name in ("区間", "車ID"):
                    continue
                old_val = old_row.get(col_name) or ""
                new_val = table.at[i, col_name] or ""
                if str(old_val) != str(new_val):
                    changed_mask.at[i, col_name] = True
                    table.at[i, col_name] = f"{new_val}　（旧: {old_val or '空欄'}）"

    # 直前の区間でランナーだった(=走者回収された)人のセルに🔴マークを付ける
    # (Excel出力側は赤枠で囲むが、st.dataframeはCSSのborderを描画しないため
    # 見た目上の代替としてマークを使う)。diffの比較には影響しないよう、
    # 修理前後の比較(上のchanged_mask算出)より後に追加する。
    for i, pickups in enumerate(pickup_cols):
        for col in pickups:
            col_name = header[col]
            table.at[i, col_name] = f"🔴 {table.at[i, col_name]}"

    def _next_run_cell_styles(_data: pd.DataFrame) -> pd.DataFrame:
        styles = pd.DataFrame("", index=table.index, columns=table.columns)
        for i, fills in enumerate(cell_fills):
            for col, color in fills.items():
                styles.iat[i, col] = f"background-color: {color}"
        return styles

    styled_table = table.style.apply(_next_run_cell_styles, axis=None)
    caption = (
        "「先行」列 = その車の乗車メンバーが次に変わる区間"
        "（直前区間のランナーを新たに乗せた行は「走者回収&◯区」）"
        "　｜　🔴 ＝直前の区間でランナーだった人(走者回収された人、Excel上は赤枠)"
    )
    if changed_mask.to_numpy().any():
        border_styles = pd.DataFrame("", index=table.index, columns=table.columns)
        for i, (label, car_id) in enumerate(row_keys):
            if car_id is None or (label, car_id) in revert_car_keys:
                continue
            for col_name in header:
                if changed_mask.at[i, col_name]:
                    border_styles.at[i, col_name] = "color: #CC3300; font-weight: bold"
        styled_table = styled_table.apply(lambda _: border_styles, axis=None)
        caption += (
            "　｜　✏️ 赤文字＝修理で内容が変わり適用されるセル"
            "　｜　（旧:...）だけ表示＝変わったが適用しないセル（内容は修理前の値のまま）"
        )

    st.dataframe(styled_table, hide_index=True, use_container_width=True, key=f"{key_prefix}_car_table")
    st.caption(caption)
    _render_section_color_legend()


def _select_source_plan(key_prefix: str):
    """表シュミの出力形式(入力データ／区間別ランナー／区間別配車の3シート)のファイルや
    スプレッドシートを選んでplan/participantsを復元する、入力元選択UI。裏シュミ・修理の
    どちらからも同じ部品を使う(key_prefixでウィジェットキーの衝突を避ける)。
    選べていなければ(None, None, None)を返す。raw_bytesは元ファイルの生のバイト列
    (「直前の表シュミ結果を使う」の場合も、表シュミが同時に保存したOUTPUT_XLSX_PATHを
    読んで返す。それも無ければNone)。修理機能が、修理前の個人別まとめの内容との比較や、
    修理対象外の内容を保持したまま書き出し直すための土台として使う。"""
    has_forward_result = bool(st.session_state.get("result"))
    has_saved_output = os.path.exists(OUTPUT_XLSX_PATH)

    source_options = []
    if has_forward_result:
        source_options.append("直前の表シュミ結果を使う")
    if has_saved_output:
        source_options.append("保存済みの表シュミ結果ファイルを使う")
    source_options.append("表シュミの結果xlsxをアップロード")
    source_options.append("表シュミの結果と同じ形式のスプレッドシートのURLを指定")

    source = st.radio("入力元", source_options, horizontal=True, key=f"{key_prefix}_source")

    plan = None
    participants = None
    raw_bytes = None

    if source == "直前の表シュミ結果を使う":
        plan = st.session_state["result"]["plan"]
        participants = st.session_state["result"]["participants"]
        if os.path.exists(OUTPUT_XLSX_PATH):
            try:
                with open(OUTPUT_XLSX_PATH, "rb") as f:
                    raw_bytes = f.read()
            except Exception:
                raw_bytes = None
    elif source == "保存済みの表シュミ結果ファイルを使う":
        st.caption(f"前回保存された結果ファイル（{OUTPUT_XLSX_PATH}）を読み込みます。")
        try:
            with open(OUTPUT_XLSX_PATH, "rb") as f:
                raw_bytes = f.read()
            participants = import_participants_from_output_xlsx(io.BytesIO(raw_bytes))
            plan = import_plan_from_output_xlsx(io.BytesIO(raw_bytes), participants)
            st.success(f"参加者 {len(participants)} 名のデータを読み込みました。")
        except Exception as e:
            st.error(f"ファイルの読み込みに失敗しました。\n\n{e}")
    elif source == "表シュミの結果と同じ形式のスプレッドシートのURLを指定":
        sheet_url = st.text_input(
            "スプレッドシートのURL（入力データ／区間別ランナー／区間別配車の3シートを含む、"
            "「リンクを知っている全員が閲覧可」に共有されたもの）",
            key=f"{key_prefix}_sheet_url",
        )
        if sheet_url:
            try:
                raw_bytes = fetch_output_format_xlsx_from_google_sheet(sheet_url).getvalue()
                participants = import_participants_from_output_xlsx(io.BytesIO(raw_bytes))
                plan = import_plan_from_output_xlsx(io.BytesIO(raw_bytes), participants)
                st.success(f"参加者 {len(participants)} 名のデータを読み込みました。")
            except Exception as e:
                st.error(f"スプレッドシートの読み込みに失敗しました。表シュミの出力形式か、共有設定を確認してください。\n\n{e}")
    else:
        uploaded = st.file_uploader(
            "表シュミが出力したxlsxファイル（入力データ／区間別ランナー／区間別配車の3シートを含むもの）",
            type=["xlsx"], key=f"{key_prefix}_uploaded_file",
        )
        if uploaded is not None:
            try:
                raw_bytes = uploaded.getvalue()
                participants = import_participants_from_output_xlsx(io.BytesIO(raw_bytes))
                plan = import_plan_from_output_xlsx(io.BytesIO(raw_bytes), participants)
                st.success(f"参加者 {len(participants)} 名のデータを読み込みました。")
            except Exception as e:
                st.error(f"ファイルの読み込みに失敗しました。表シュミの出力形式のxlsxか確認してください。\n\n{e}")

    return plan, participants, raw_bytes


def _load_old_individual_summary_by_name(raw_bytes):
    """修理前の「個人別まとめ」シートに残っていたキャッシュ済みの表示値(元ファイルが
    実際にGoogle Sheets/Excelで開かれた際に保存されたもの)を、名前をキーに読み込む。
    元ファイルにキャッシュが無い(作成後に一度も開かれていない等)場合はNoneを返し、
    比較できないことを示す。"""
    if not raw_bytes:
        return None
    try:
        wb = load_workbook(io.BytesIO(raw_bytes), data_only=True)
    except Exception:
        return None
    if "個人別まとめ" not in wb.sheetnames:
        return None
    ws = wb["個人別まとめ"]
    header = [c.value for c in ws[1]]
    if "名前" not in header:
        return None
    name_col = header.index("名前") + 1
    # 区間ラベル以外の列(「凡例」の見出しなど、同じ行範囲に横並びで書かれている
    # 凡例欄)を取り違えて「キャッシュあり」と誤判定しないよう、比較対象は
    # _diff_individual_summaryが実際に見る区間ラベルだけに絞る。
    section_labels = set(section_label(s) for s in list(range(1, 11)) + [11])

    old_by_name: Dict[str, Dict[str, str]] = {}
    any_value_found = False
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=name_col).value
        if not name:
            continue
        row_vals = {}
        for c, label in enumerate(header, start=1):
            if label not in section_labels:
                continue
            v = ws.cell(row=r, column=c).value
            if v:
                any_value_found = True
            row_vals[label] = v
        old_by_name[name] = row_vals

    if not any_value_found:
        return None
    return old_by_name


def _diff_individual_summary(old_by_name, participants, new_summary):
    """修理前の「個人別まとめ」の表示値(_load_old_individual_summary_by_nameの結果)と、
    修理後の内容をセル単位で突き合わせ、変わった箇所だけを返す。old_by_nameがNone
    (比較できない)ならNoneを返す。"""
    if old_by_name is None:
        return None
    labels = [section_label(s) for s in list(range(1, 11)) + [11]]
    diffs = []
    for pid, p in participants.items():
        old_row = old_by_name.get(p.name, {})
        new_row = new_summary.get(pid, {})
        for label in labels:
            old_val = old_row.get(label) or ""
            new_val = new_row.get(label) or ""
            if old_val != new_val:
                diffs.append({
                    "適用する": True, "名前": p.name, "区間": label,
                    "修理前": old_val or "(空欄)", "修理後": new_val or "(空欄)",
                })
    return diffs


def _load_old_cars_table(raw_bytes):
    """修理前の「区間別配車」シートの表示値を、(区間ラベル, 車ID)をキーに読み込む。
    このシートは元々数式を使わない(_write_cars_sheet参照)ため、個人別まとめと違い
    「一度もExcelで開かれていないとキャッシュが無い」という制約が無く、常に実際の
    値を読める。raw_bytesが無い/シート構成が想定と違う場合はNoneを返す。"""
    if not raw_bytes:
        return None
    try:
        wb = load_workbook(io.BytesIO(raw_bytes), data_only=True)
    except Exception:
        return None
    if "区間別配車" not in wb.sheetnames:
        return None
    ws = wb["区間別配車"]
    header = [c.value for c in ws[1]]
    if "車ID" not in header or "区間" not in header:
        return None
    car_id_col = header.index("車ID") + 1

    title_row = None
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.startswith("■"):
            title_row = r
            break
    last_row = (title_row - 1) if title_row else ws.max_row

    old_by_key: Dict[tuple, Dict[str, object]] = {}
    current_label = None
    for r in range(2, last_row + 1):
        label_cell = ws.cell(row=r, column=1).value
        if label_cell:
            current_label = label_cell
        car_id = ws.cell(row=r, column=car_id_col).value
        if current_label is None or not car_id:
            continue
        row_vals = {}
        for c, col_name in enumerate(header, start=1):
            if not col_name or col_name in ("区間", "車ID"):
                continue
            row_vals[col_name] = ws.cell(row=r, column=c).value
        old_by_key[(current_label, str(car_id))] = row_vals
    return old_by_key


def _cars_diff_rows(raw_bytes, plan, participants):
    """区間別配車で修理/裏シュミにより実際に変わる行だけを、UIで選択できる一覧として
    返す(「適用する」列付き)。raw_bytesが無い、ファイルが読めない、またはシートの
    見出しが想定と違って比較できない場合はNoneを返す(区別のため、差分が無い場合は
    空リストを返す)。"""
    if not raw_bytes:
        return None
    try:
        wb = load_workbook(io.BytesIO(raw_bytes))
    except Exception:
        return None
    if "区間別配車" not in wb.sheetnames:
        return None
    changes = describe_car_row_changes(wb["区間別配車"], plan, participants)
    if changes is None:
        return None
    rows = []
    for c in changes:
        if not c["changed"]:
            continue
        parts = []
        if (c["old_advance"] or None) != (c["new_advance"] or None):
            parts.append(f"先行: {c['old_advance'] or '(空欄)'} → {c['new_advance'] or '(空欄)'}")
        if set(c["old_names"]) != set(c["new_names"]):
            old_text = "、".join(c["old_names"]) or "(空欄)"
            new_text = "、".join(c["new_names"]) or "(空欄)"
            parts.append(f"乗車者: {old_text} → {new_text}")
        rows.append({
            "適用する": True, "区間": c["label"], "車ID": c["car_id"],
            "変更内容": " ／ ".join(parts),
        })
    return rows


def _apply_individual_diff_selection(new_summary, old_individual, participants, apply_individual):
    """個人別まとめの差分ごとの適用フラグ(apply_individual、(名前,区間)→bool)を見て、
    適用しないと選ばれたセルだけ修理前の値に戻したsummaryを返す(new_summary自体は
    書き換えない)。"""
    summary = {pid: dict(v) for pid, v in new_summary.items()}
    if old_individual is None:
        return summary
    name_to_pid = {p.name: pid for pid, p in participants.items()}
    for (name, label), apply in apply_individual.items():
        if apply:
            continue
        pid = name_to_pid.get(name)
        if pid is None:
            continue
        old_val = old_individual.get(name, {}).get(label) or ""
        summary[pid][label] = old_val
    return summary


def _revert_car_keys_from_selection(apply_cars):
    """区間別配車の差分ごとの適用フラグ(apply_cars、(区間,車ID)→bool)から、適用しないと
    選ばれた(区間,車ID)のsetを返す。"""
    return {key for key, apply in apply_cars.items() if not apply}



def _render_repair_section():
    """修理: 表シュミ・裏シュミの出力と同じ形式のファイル/スプレッドシートを読み込み、
    人員の交代や入力データとの不整合の調整は一切せず、同じ4シート構成のExcelを
    数式を使わずに書き出し直す。区間別配車のセルの色分け・「先行」列は裏シュミのUIの
    区間別配車タブと同じ(compute_car_table_overridesを共用)。"""
    st.subheader("🔧 修理（数式を使わず同じ形式で書き出し直す）")
    st.caption(
        "区間別配車・区間別ランナーをドラッグ移動などで手動編集すると、個人別まとめの数式が"
        "壊れて「重複」の誤検知や#N/Aが出ることがあります。この機能は人員の交代や入力データとの"
        "不整合の調整は一切行わず、今の内容をそのまま、数式を使わない同じ形式のExcelファイルとして"
        "書き出し直すだけです。"
    )

    plan, participants, raw_bytes = _select_source_plan("repair")
    if plan is None or participants is None:
        st.info("修理するには、上で入力元を指定してください。")
        return

    if st.button("🔧 修理を実行", type="primary", key="repair_run"):
        if not raw_bytes:
            st.error("修理の土台となる元ファイルが見つかりませんでした。もう一度入力元を選び直してください。")
            return
        new_summary_full = compute_individual_summary_static(plan, participants)
        old_individual = _load_old_individual_summary_by_name(raw_bytes)
        old_cars = _load_old_cars_table(raw_bytes)
        individual_diffs = _diff_individual_summary(old_individual, participants, new_summary_full)
        cars_diffs = _cars_diff_rows(raw_bytes, plan, participants)
        st.session_state.repair_result = {
            "plan": plan, "participants": participants, "raw_bytes": raw_bytes,
            "new_summary_full": new_summary_full, "old_individual": old_individual, "old_cars": old_cars,
            "individual_diffs": individual_diffs, "cars_diffs": cars_diffs,
        }
        # 「適用する/しない」の状態は、この2つの辞書を単一の根拠(source of truth)として
        # 持つ。チェックボックス表・一括ボタン・プレビュー表での選択操作は、いずれも
        # この辞書を更新してからrepair_revisionを進め、ウィジェットのキーを変えて
        # 強制的に再構築させることで反映する(st.data_editor自身の編集履歴と
        # 食い違わないようにするため)。
        st.session_state.repair_apply_individual = {
            (row["名前"], row["区間"]): True for row in (individual_diffs or [])
        }
        st.session_state.repair_apply_cars = {
            (row["区間"], row["車ID"]): True for row in (cars_diffs or [])
        }
        st.session_state.repair_revision = 0

    repair_result = st.session_state.get("repair_result")
    if not repair_result:
        return

    r_plan = repair_result["plan"]
    r_participants = repair_result["participants"]
    r_raw_bytes = repair_result["raw_bytes"]
    individual_diffs = repair_result.get("individual_diffs")
    cars_diffs = repair_result.get("cars_diffs")
    apply_individual = st.session_state.setdefault("repair_apply_individual", {})
    apply_cars = st.session_state.setdefault("repair_apply_cars", {})
    revision = st.session_state.get("repair_revision", 0)

    def _bump_revision():
        st.session_state.repair_revision = st.session_state.get("repair_revision", 0) + 1
        st.rerun()

    st.markdown("#### 修理で変わる内容の確認")
    st.caption(
        "チェックを外した行、または下の個人別まとめ・区間別配車の表で選んで「適用しない」に"
        "した箇所は、実際の出力ファイルでは修理前の内容のまま保たれます。"
    )

    if individual_diffs is None:
        st.caption("ℹ️ 元ファイルの個人別まとめに保存済みの表示内容が見つからず、修理前との比較はできませんでした。")
    elif len(individual_diffs) == 0:
        st.caption("個人別まとめの内容に、修理前との差分はありませんでした。")
    else:
        st.markdown(f"**個人別まとめで変更された箇所（{len(individual_diffs)}件）**")
        bcol1, bcol2 = st.columns(2)
        if bcol1.button("すべて適用する", key="repair_individual_select_all"):
            for k in apply_individual:
                apply_individual[k] = True
            _bump_revision()
        if bcol2.button("すべて適用しない", key="repair_individual_select_none"):
            for k in apply_individual:
                apply_individual[k] = False
            _bump_revision()

        individual_df = pd.DataFrame([
            {**row, "適用する": apply_individual.get((row["名前"], row["区間"]), True)}
            for row in individual_diffs
        ])
        edited_individual_df = st.data_editor(
            individual_df, hide_index=True, use_container_width=True,
            key=f"repair_individual_diff_editor_{revision}",
            disabled=["名前", "区間", "修理前", "修理後"],
            column_config={"適用する": st.column_config.CheckboxColumn(default=True)},
        )
        for _, r in edited_individual_df.iterrows():
            apply_individual[(r["名前"], r["区間"])] = bool(r["適用する"])

    final_summary = _apply_individual_diff_selection(
        repair_result["new_summary_full"], repair_result["old_individual"], r_participants, apply_individual)

    if individual_diffs:
        unapplied_individual_keys = {k for k, v in apply_individual.items() if not v}
        st.markdown("**👤 個人別まとめ（区間ごとの役割）**")
        _render_individual_summary_table(
            r_plan, r_participants, key_prefix="repair", summary=repair_result["new_summary_full"],
            old_summary=repair_result.get("old_individual"), unapplied_keys=unapplied_individual_keys,
        )

    if cars_diffs is None:
        st.caption("ℹ️ 元ファイルの区間別配車シートの形式が確認できず、修理前との比較はできませんでした。")
    elif len(cars_diffs) == 0:
        st.caption("区間別配車の内容に、修理前との差分はありませんでした。")
    else:
        st.markdown(f"**区間別配車で変更された箇所（{len(cars_diffs)}件）**")
        bcol1, bcol2 = st.columns(2)
        if bcol1.button("すべて適用する", key="repair_cars_select_all"):
            for k in apply_cars:
                apply_cars[k] = True
            _bump_revision()
        if bcol2.button("すべて適用しない", key="repair_cars_select_none"):
            for k in apply_cars:
                apply_cars[k] = False
            _bump_revision()

        cars_df = pd.DataFrame([
            {**row, "適用する": apply_cars.get((row["区間"], row["車ID"]), True)}
            for row in cars_diffs
        ])
        edited_cars_df = st.data_editor(
            cars_df, hide_index=True, use_container_width=True,
            key=f"repair_cars_diff_editor_{revision}",
            disabled=["区間", "車ID", "変更内容"],
            column_config={"適用する": st.column_config.CheckboxColumn(default=True)},
        )
        for _, r in edited_cars_df.iterrows():
            apply_cars[(r["区間"], r["車ID"])] = bool(r["適用する"])

    revert_car_keys = _revert_car_keys_from_selection(apply_cars)

    if cars_diffs:
        st.markdown("**🚘 区間別配車（Excelの「区間別配車」シートと同じ形式）**")
        _render_car_assignment_table(
            r_plan, r_participants, key_prefix="repair",
            old_table=repair_result.get("old_cars"), revert_car_keys=revert_car_keys,
        )

    write_repaired_xlsx(
        io.BytesIO(r_raw_bytes), r_plan, r_participants, REPAIR_OUTPUT_XLSX_PATH,
        revert_car_keys=revert_car_keys, individual_summary=final_summary,
    )
    st.success("修理済みのファイルを作成しました（選択内容を反映済み）。")

    if os.path.exists(REPAIR_OUTPUT_XLSX_PATH):
        with open(REPAIR_OUTPUT_XLSX_PATH, "rb") as f:
            st.download_button(
                "📊 修理済みExcelをダウンロード", f,
                file_name=os.path.basename(REPAIR_OUTPUT_XLSX_PATH),
                mime=XLSX_MIME, type="primary", key="repair_download",
            )


def _render_back_sim_section():
    """裏シュミ(追加条件での再調整)。表シュミを実行していなくても、表シュミ出力形式の
    xlsxファイルを直接アップロードすれば単独で使える。"""
    st.subheader("🔧 裏シュミ（追加条件で再調整）")
    st.caption(
        "既存の配車結果を土台に、「できるだけ一緒にしたい」「離したい」「運転区間数の範囲」を"
        "追加のお願いとして加え、既存の割り当てをできるだけ変えずに調整します。"
        "免許・定員などの必須条件や、表シュミ側の希望(特に「特に走りたい区間」)はここでの新しい条件より優先されます。"
    )

    plan, participants, raw_bytes = _select_source_plan("back")
    if plan is None or participants is None:
        st.info("裏シュミを実行するには、上で入力元を指定してください。")
        return

    _, unmet_priority = check_priority_sections(plan, participants)
    if unmet_priority:
        names = "、".join(
            f"{participants[p].name}（{s}区）"
            for p, s in sorted(unmet_priority, key=lambda x: x[1])
        )
        st.warning(f"⚠️ 「特に走りたい区間」が未充足の人がいます: {names}　→ 裏シュミ実行時に優先的に直します。")

    # 学年順(同学年内は名前順)に並べる。一括選択(学年ボタン)にも使う。
    participants_sorted = sorted(participants.values(), key=lambda p: (p.grade, p.name))
    participant_names = [p.name for p in participants_sorted]
    driver_names = [p.name for p in participants_sorted if p.can_drive]
    names_by_grade: Dict[int, List[str]] = {}
    for p in participants_sorted:
        names_by_grade.setdefault(p.grade, []).append(p.name)
    grade_options = sorted(names_by_grade.keys())

    def _add_grade_members(multiselect_key, grade, exclude_name):
        current = list(st.session_state.get(multiselect_key, []))
        additions = [n for n in names_by_grade.get(grade, []) if n != exclude_name and n not in current]
        st.session_state[multiselect_key] = current + additions

    def _remove_group(count_key: str, prefix: str, fields: List[str], idx: int, total: int):
        """グループidxを削除し、それ以降のグループの入力値を1つずつ前に詰める。"""
        for j in range(idx, total - 1):
            for field in fields:
                src, dst = f"{prefix}_{field}_{j + 1}", f"{prefix}_{field}_{j}"
                if src in st.session_state:
                    st.session_state[dst] = st.session_state[src]
                elif dst in st.session_state:
                    del st.session_state[dst]
        last = total - 1
        for field in fields:
            st.session_state.pop(f"{prefix}_{field}_{last}", None)
        st.session_state[count_key] = total - 1

    NO_SELECTION = "（未選択）"
    anchor_options = [NO_SELECTION] + participant_names

    with st.expander("裏シュミの条件を設定", expanded=not st.session_state.get("back_result")):
        st.markdown("**できるだけ一緒にしたい**")
        n_together = st.number_input(
            "グループ数", min_value=0, max_value=10, step=1, key="n_together",
        )
        together_pairs = []
        together_or_groups = []
        for i in range(int(n_together)):
            c1, c2 = st.columns(2)
            anchor = c1.selectbox("基準となる人", anchor_options, key=f"together_anchor_{i}")
            companions = c2.multiselect(
                "一緒にしたい相手（複数選択可）",
                [n for n in participant_names if n != anchor],
                key=f"together_companions_{i}",
            )
            gcol1, gcol2 = st.columns([3, 1])
            grade_pick = gcol1.selectbox(
                "学年で一括追加", grade_options, key=f"together_grade_pick_{i}", label_visibility="collapsed",
            )
            gcol2.button(
                f"{grade_pick}年を追加", key=f"together_grade_btn_{i}",
                on_click=_add_grade_members, args=(f"together_companions_{i}", grade_pick, anchor),
            )
            mode = st.radio(
                "複数選んだ場合の扱い",
                ["それぞれとできるだけ一緒に（AND）", "誰か一人とでも一緒ならOK（OR）"],
                key=f"together_mode_{i}", horizontal=True,
            )
            st.button(
                "🗑️ このグループを削除", key=f"together_remove_{i}",
                on_click=_remove_group, args=("n_together", "together", ["anchor", "companions", "mode"], i, int(n_together)),
            )
            if anchor != NO_SELECTION and companions:
                if mode.startswith("誰か"):
                    together_or_groups.append((anchor, companions))
                else:
                    together_pairs += [(anchor, comp) for comp in companions]
            st.divider()

        st.markdown("**できるだけ離したい**")
        n_apart = st.number_input(
            "グループ数", min_value=0, max_value=10, step=1, key="n_apart",
        )
        apart_pairs = []
        for i in range(int(n_apart)):
            c1, c2 = st.columns(2)
            anchor = c1.selectbox("基準となる人", anchor_options, key=f"apart_anchor_{i}")
            companions = c2.multiselect(
                "離したい相手（複数選択可）",
                [n for n in participant_names if n != anchor],
                key=f"apart_companions_{i}",
            )
            gcol1, gcol2 = st.columns([3, 1])
            grade_pick = gcol1.selectbox(
                "学年で一括追加", grade_options, key=f"apart_grade_pick_{i}", label_visibility="collapsed",
            )
            gcol2.button(
                f"{grade_pick}年を追加", key=f"apart_grade_btn_{i}",
                on_click=_add_grade_members, args=(f"apart_companions_{i}", grade_pick, anchor),
            )
            st.button(
                "🗑️ このグループを削除", key=f"apart_remove_{i}",
                on_click=_remove_group, args=("n_apart", "apart", ["anchor", "companions"], i, int(n_apart)),
            )
            if anchor != NO_SELECTION and companions:
                apart_pairs += [(anchor, comp) for comp in companions]
            st.divider()

        st.markdown("**運転区間数の範囲**")
        n_driver_ranges = st.number_input(
            "指定する人数", min_value=0, max_value=20, value=0, step=1, key="n_driver_ranges",
        )
        driver_ranges = {}
        DRIVER_RANGE_NO_LIMIT = 11
        for i in range(int(n_driver_ranges)):
            c1, c2, c3 = st.columns(3)
            person = c1.selectbox("対象者", driver_names, key=f"driver_range_person_{i}")
            min_r = c2.number_input(
                "最少運転区間数", min_value=0, max_value=DRIVER_RANGE_NO_LIMIT, value=0, step=1,
                key=f"driver_range_min_{i}",
            )
            max_r = c3.number_input(
                "最大運転区間数（11で無制限）", min_value=0, max_value=DRIVER_RANGE_NO_LIMIT,
                value=DRIVER_RANGE_NO_LIMIT, step=1, key=f"driver_range_max_{i}",
            )
            driver_ranges[person] = (int(min_r), None if int(max_r) >= DRIVER_RANGE_NO_LIMIT else int(max_r))

        back_time_limit = st.number_input(
            "裏シュミの計算制限時間（秒/フェーズ）", min_value=10, max_value=600, value=60, step=10,
            key="back_time_limit",
        )

        back_output_format = st.radio(
            "出力するExcelの形式", ["関数組み込みVer（バグが起こる場合があります）", "関数なしVer"],
            horizontal=True, key="back_output_format",
            help="関数なしVerでは、区間別配車・区間別ランナーのうち裏シュミで実際に変わった箇所だけを"
                 "書き換え、それ以外(メモ・運転手一覧など)は入力元ファイルのまま保持します。",
        )
        back_use_formulas = back_output_format.startswith("関数組み込み")

        has_any_condition = bool(together_pairs or together_or_groups or apart_pairs or driver_ranges)
        needs_optimization = has_any_condition or bool(unmet_priority)
        if has_any_condition:
            run_help = None
        elif unmet_priority:
            run_help = "条件は指定されていませんが、「特に走りたい区間」の未充足を直すために再計算します。"
        else:
            run_help = "条件が指定されていないため、再計算せずそのまま結果を表示します（区間別配車の表の確認などに）。"
        run_back = st.button("裏シュミを実行", type="primary", help=run_help)

    def _write_back_output(back_plan):
        if back_use_formulas:
            write_plan_xlsx(back_plan, participants, BACK_OUTPUT_XLSX_PATH, use_formulas=True)
        elif raw_bytes:
            write_updated_output_xlsx(io.BytesIO(raw_bytes), back_plan, participants, BACK_OUTPUT_XLSX_PATH)
        else:
            # 元ファイルが無い(入力元が「直前の表シュミ結果を使う」で保存前など)場合は、
            # 修理対象外の内容を保持できないが、新形式そのものは出力できるようにする。
            write_plan_xlsx(back_plan, participants, BACK_OUTPUT_XLSX_PATH, use_formulas=False)

    if run_back:
        back_log_buf = io.StringIO()
        config = BackSimConfig(
            together_pairs=together_pairs, apart_pairs=apart_pairs, driver_ranges=driver_ranges,
            together_or_groups=together_or_groups,
        )
        if not needs_optimization:
            back_log_buf.write("条件が指定されておらず、「特に走りたい区間」も充足済みのため、計算をせずに読み込んだ結果をそのまま表示します。\n")
            back_plan = plan
            _write_back_output(back_plan)
            condition_summary = summarize_conditions(back_plan, participants, config)
        else:
            with st.spinner("裏シュミを計算中..."):
                try:
                    with contextlib.redirect_stdout(back_log_buf):
                        back_plan = run_back_sim(plan, participants, config, time_limit=back_time_limit)
                        _write_back_output(back_plan)
                        condition_summary = summarize_conditions(back_plan, participants, config)
                except Exception as e:
                    st.error(f"裏シュミの計算に失敗しました。\n\n{e}")
                    with st.expander("詳細ログ"):
                        st.text(back_log_buf.getvalue())
                    st.stop()

        together_map: Dict[str, set] = {}
        for a, b in together_pairs:
            together_map.setdefault(a, set()).add(b)
            together_map.setdefault(b, set()).add(a)
        for or_anchor, or_group in together_or_groups:
            for g in or_group:
                together_map.setdefault(or_anchor, set()).add(g)
                together_map.setdefault(g, set()).add(or_anchor)

        apart_map: Dict[str, set] = {}
        for a, b in apart_pairs:
            apart_map.setdefault(a, set()).add(b)
            apart_map.setdefault(b, set()).add(a)

        st.session_state.back_result = {
            "plan": back_plan,
            "participants": participants,
            "log": back_log_buf.getvalue(),
            "summary": condition_summary,
            "highlight": {
                "together_names": set(together_map.keys()),
                "apart_names": set(apart_map.keys()),
                "driver_range_names": set(driver_ranges.keys()),
                "together_map": together_map,
                "apart_map": apart_map,
            },
        }

    back_result = st.session_state.get("back_result")
    if back_result:
        st.markdown("#### 裏シュミの結果")

        summary = back_result.get("summary", {})
        has_any_summary = any(
            summary.get(k) for k in ("together", "together_or", "apart", "driver_ranges", "priority_sections")
        )
        if has_any_summary:
            st.markdown("**条件の充足状況**")
            if summary.get("priority_sections"):
                st.caption("特に走りたい区間（表シュミ側の希望。裏シュミの新条件より優先して充足を試みます）")
                st.dataframe(
                    pd.DataFrame([
                        {"名前": d["name"], "区間": f"{d['section_id']}区", "充足": "✅" if d["satisfied"] else "❌"}
                        for d in summary["priority_sections"]
                    ]),
                    hide_index=True, use_container_width=True, key="back_summary_priority_sections",
                )
            if summary.get("together"):
                st.caption("できるだけ一緒にしたい（同乗できた区間数）")
                st.dataframe(
                    pd.DataFrame([
                        {"人物A": t["a"], "人物B": t["b"], "同乗できた区間数": t["shared_sections"]}
                        for t in summary["together"]
                    ]),
                    hide_index=True, use_container_width=True, key="back_summary_together",
                )
            if summary.get("together_or"):
                st.caption("できるだけ一緒にしたい（候補のうち誰か一人とでも同乗できた区間数）")
                st.dataframe(
                    pd.DataFrame([
                        {
                            "基準となる人": t["anchor"],
                            "候補": "、".join(t["group"]),
                            "満たせた区間数": t["satisfied_sections"],
                        }
                        for t in summary["together_or"]
                    ]),
                    hide_index=True, use_container_width=True, key="back_summary_together_or",
                )
            if summary.get("apart"):
                st.caption("できるだけ離したい（同乗してしまった区間数。0が理想）")
                st.dataframe(
                    pd.DataFrame([
                        {"人物A": t["a"], "人物B": t["b"], "同乗してしまった区間数": t["shared_sections"]}
                        for t in summary["apart"]
                    ]),
                    hide_index=True, use_container_width=True, key="back_summary_apart",
                )
            if summary.get("driver_ranges"):
                st.caption("運転区間数の範囲")
                st.dataframe(
                    pd.DataFrame([
                        {
                            "名前": d["name"],
                            "指定範囲": f"{d['min']}〜{d['max'] if d['max'] is not None else '無制限'}",
                            "実際の運転区間数": d["actual"],
                            "達成": "✅" if d["satisfied"] else "❌",
                        }
                        for d in summary["driver_ranges"]
                    ]),
                    hide_index=True, use_container_width=True, key="back_summary_driver_ranges",
                )

        _render_plan_result(
            back_result["plan"], back_result["participants"], BACK_OUTPUT_XLSX_PATH, key_prefix="back",
            download_label="📊 裏シュミの結果Excelをダウンロード",
            highlight=back_result.get("highlight"),
        )
        with st.expander("裏シュミの詳細ログ"):
            st.text(back_result["log"])


# =============================================================================
# 画面構成
# =============================================================================

st.set_page_config(page_title="箱根駅伝配車シミュレーター (CP-SAT版)", page_icon="🚗")
st.title("🚗 箱根駅伝配車シミュレーター")
st.caption(f"ver {APP_VERSION}　(計算エンジン: OR-Tools CP-SAT)")
_hidden_dot("dot_1", "hidden_dot_1")
st.divider()

# --- 表シュミ ---
input_format = st.radio(
    "入力データの形式",
    ["Googleフォームの回答", "独自フォーマット"],
    horizontal=True,
)

if input_format == "Googleフォームの回答":
    st.caption("Googleフォームの回答が集まったスプレッドシートのURLを貼り付けてください。")
else:
    st.caption("所定のフォーマット（名前・1〜10・運転・大・山・宿泊・学年・希望区間数・離脱区間）のスプレッドシートのURLを貼り付けてください。")

url = st.text_input(
    "スプレッドシートのURL",
    placeholder="https://docs.google.com/spreadsheets/d/...",
)

auth_mode = st.radio(
    "スプレッドシートへのアクセス方法",
    ["認証なし（読み込み専用・推奨）", "サービスアカウント認証（読み込み＋書き戻し）"],
    horizontal=True,
    help=(
        "「認証なし」は、スプレッドシートが「リンクを知っている全員が閲覧可」に共有されていれば"
        "Google Cloudの設定が一切不要です。結果は画面表示とExcelダウンロードで受け取ります。\n\n"
        "「サービスアカウント認証」は結果をスプレッドシートに書き戻せますが、"
        "credentials.json（Google Cloudのサービスアカウントキー）が必要です。"
    ),
)
use_auth = auth_mode.startswith("サービスアカウント")
if use_auth and not os.path.exists(CREDENTIALS_PATH):
    st.warning(f"credentials.json が見つかりません（{os.path.abspath(CREDENTIALS_PATH)}）。認証なしモードに切り替えてください。")
_hidden_dot("dot_3", "hidden_dot_3")

st.subheader("車両スロット設定")
col_l, col_n = st.columns(2)
n_large_slots = col_l.number_input(
    "大型車（8人乗り）スロット数", min_value=0, max_value=len(LARGE_CAR_IDS), value=0, step=1
)
n_normal_slots = col_n.number_input(
    "普通車（4人乗り）スロット数", min_value=0, max_value=len(NORMAL_CAR_IDS), value=len(NORMAL_CAR_IDS), step=1
)
active_car_ids = LARGE_CAR_IDS[:n_large_slots] + NORMAL_CAR_IDS[:n_normal_slots]
st.caption(f"最大定員: 大型 {n_large_slots}台×8人 + 普通 {n_normal_slots}台×4人 = {n_large_slots*8 + n_normal_slots*4}人")
_hidden_dot("dot_2", "hidden_dot_2")

NO_LIMIT = 99  # 「無制限」を表す便宜上の大きな値(実際の参加人数を超えるので事実上上限なしになる)


def _apply_bulk_limits():
    bulk_min = st.session_state.get("bulk_limit_min", 1)
    bulk_max = st.session_state.get("bulk_limit_max", NO_LIMIT)
    for s in range(1, 11):
        st.session_state[f"limit_min_{s}"] = int(bulk_min)
        st.session_state[f"limit_max_{s}"] = int(bulk_max)


runner_limits = {}
with st.expander("区間ごとの人数制限（任意）"):
    st.caption(
        "各区間を走る人数の下限・上限を指定できます。最大人数は99のままにすると"
        "実質「上限なし」として扱われます。指定しない区間は最少1人・上限なしの既定値になります。"
    )

    st.markdown("**一括設定**（下の値を1〜10区すべてに適用します）")
    bulk_col_min, bulk_col_max = st.columns(2)
    bulk_col_min.number_input(
        "最少人数（一括）", min_value=0, max_value=NO_LIMIT, value=1, step=1, key="bulk_limit_min",
    )
    bulk_col_max.number_input(
        "最大人数（一括）", min_value=0, max_value=NO_LIMIT, value=NO_LIMIT, step=1, key="bulk_limit_max",
    )
    st.button("↑ 1〜10区すべてに適用", on_click=_apply_bulk_limits, key="apply_bulk_limits_btn")

    st.divider()

    col_label, col_min, col_max = st.columns([1, 2, 2])
    col_min.markdown("**最少人数**")
    col_max.markdown("**最大人数**")
    for s in range(1, 11):
        st.session_state.setdefault(f"limit_min_{s}", 1)
        st.session_state.setdefault(f"limit_max_{s}", NO_LIMIT)
        col_label, col_min, col_max = st.columns([1, 2, 2])
        col_label.markdown(f"{s}区")
        min_n = col_min.number_input(
            f"{s}区の最少人数", min_value=0, max_value=NO_LIMIT, step=1,
            key=f"limit_min_{s}", label_visibility="collapsed",
        )
        max_n = col_max.number_input(
            f"{s}区の最大人数", min_value=0, max_value=NO_LIMIT, step=1,
            key=f"limit_max_{s}", label_visibility="collapsed",
        )
        runner_limits[s] = (int(min_n), None if int(max_n) >= NO_LIMIT else int(max_n))

time_limit = st.number_input(
    "1ブロックあたりの計算制限時間（秒）", min_value=10, max_value=600, value=int(DEFAULT_TIME_LIMIT), step=10,
    help="大型車が少ない構成では最適解を見つけるのに数分かかることがあります。",
)

output_format = st.radio(
    "出力するExcelの形式", ["関数組み込みVer（バグが起こる場合があります）", "関数なしVer"],
    horizontal=True, key="fwd_output_format",
    help="関数組み込みVerは区間別配車・区間別ランナーを手動編集すると個人別まとめが自動追従しますが、"
         "セルのドラッグ移動などで数式が壊れることがあります。関数なしVerは自動追従しない代わりに壊れません。",
)
use_formulas = output_format.startswith("関数組み込み")

can_write_back = use_auth and os.path.exists(CREDENTIALS_PATH)
write_back = st.checkbox(
    "結果をスプレッドシートに書き戻す（「入力データ」「配車結果」シートを作成/上書き）",
    value=can_write_back,
    disabled=not can_write_back,
    help="「サービスアカウント認証」を選び、credentials.jsonが置かれている場合のみ有効になります。" if not can_write_back else None,
)

if st.button("シミュレーション実行", type="primary", disabled=not url or not active_car_ids):
    # 計算結果はすべてsession_stateに保存する。
    # (Excelダウンロードボタンを押すとStreamlitはスクリプト全体を再実行するが、
    #  そのときst.button()はFalseに戻るため、結果をsession_state以外に置いていると
    #  再実行のたびに画面から消えてしまう。)
    log_buf = io.StringIO()

    with st.spinner("データを読み込み中..."):
        try:
            with contextlib.redirect_stdout(log_buf):
                if use_auth:
                    if input_format == "Googleフォームの回答":
                        participants = load_participants_from_form_sheet(url, CREDENTIALS_PATH)
                    else:
                        participants = load_participants_from_sheet(url, CREDENTIALS_PATH)
                else:
                    if input_format == "Googleフォームの回答":
                        participants = load_participants_from_public_form_sheet(url)
                    else:
                        participants = load_participants_from_public_sheet(url)
        except Exception as e:
            st.error(f"スプレッドシートの読み込みに失敗しました。\n\n{e}")
            st.stop()

    n_participants = len(participants)
    input_warnings = validate_participants(participants)

    with st.spinner(f"配車を計算中（大型車が少ない構成では最大{time_limit*2:.0f}秒程度かかる場合があります）..."):
        try:
            with contextlib.redirect_stdout(log_buf):
                plan = generate_full_plan_cpsat(
                    participants,
                    active_car_ids=active_car_ids,
                    time_limit=time_limit,
                    output_path=OUTPUT_XLSX_PATH,
                    runner_limits=runner_limits,
                    use_formulas=use_formulas,
                )
        except Exception as e:
            st.error(f"計算に失敗しました。\n\n{e}")
            with st.expander("詳細ログ"):
                st.text(log_buf.getvalue())
            st.stop()

    st.session_state.result = {
        "plan": plan,
        "participants": participants,
        "n_participants": n_participants,
        "input_warnings": input_warnings,
        "log": log_buf.getvalue(),
        "write_back_requested": write_back,
        "write_back_done": False,
        "write_back_url": None,
        "write_back_error": None,
        "url": url,
    }

# --- 結果表示ブロック: session_stateに結果がある限り、再実行のたびに描画する ---
result = st.session_state.get("result")
if result:
    plan = result["plan"]
    participants = result["participants"]

    st.info(f"参加者 {result['n_participants']} 名のデータを読み込みました。")

    if result["input_warnings"]:
        with st.expander(f"⚠️ 入力データに {len(result['input_warnings'])} 件の問題があります"):
            for w in result["input_warnings"]:
                st.warning(w)

    _render_plan_result(
        plan, participants, OUTPUT_XLSX_PATH, key_prefix="fwd",
        download_label="📊 結果Excelをダウンロード（入力データ／区間別ランナー／区間別配車／個人別まとめの4シート）",
    )

    if result["write_back_requested"] and not result["write_back_done"]:
        with st.spinner("スプレッドシートに結果を書き込み中..."):
            try:
                with contextlib.redirect_stdout(io.StringIO()):
                    result["write_back_url"] = save_plan_to_sheet(plan, participants, result["url"], CREDENTIALS_PATH)
            except Exception as e:
                result["write_back_error"] = str(e)
            result["write_back_done"] = True

    if result["write_back_url"]:
        st.link_button("📄 結果のスプレッドシートを開く（入力データ／配車結果シート）", result["write_back_url"], type="primary")
    if result["write_back_error"]:
        st.error(f"スプレッドシートへの書き出しに失敗しました。\n\n{result['write_back_error']}")

    with st.expander("詳細ログ"):
        st.text(result["log"])

st.divider()
_render_repair_section()

# --- 裏シュミへの入口(ページ最後のドット) ---
# 表シュミの結果はここに来るまでのスクリプト実行で既に確定しているので、そのまま使える。
st.divider()
_hidden_dot("dot_4", "hidden_dot_4")

if st.session_state.get("show_back_sim"):
    st.divider()
    _render_back_sim_section()
