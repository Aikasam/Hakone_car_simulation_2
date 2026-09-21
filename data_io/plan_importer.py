"""表シュミ(output_writer.write_plan_xlsxで書き出したxlsx)を読み戻して、
Participant/SectionState/CarStateを再構築する(裏シュミの入力用)。

入力データシートの宿泊/運転/大/山/希望区間の各フラグはセルの値ではなく塗り色で
表現されているため(_write_input_sheet参照)、色から読み戻す。区間別配車・区間別
ランナーの名前は表記ゆれ(括弧書きの注記・全角/半角スペース)を吸収してから
参加者名と突き合わせる(このセッションで個人別まとめの数式に使ってきたのと
同じ正規化ロジックを流用する)。
"""

import io
import os
import re
import sys
from typing import Dict, List, Optional

import requests
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from models import CarState, Participant, SectionState
from logic.car_pool import CAR_TYPE, section_label

FILL_WHITE = "FFFFFFFF"
FILL_BLACK = "FF000000"
FILL_RED = "FFFF0000"

INPUT_HEADER = ["名前", "学年", "宿泊", "運転", "大", "山",
                "1", "2", "3", "4", "5", "6", "7", "8", "9", "10",
                "希望区間数", "離脱区間", "特に走りたい区間"]


def fetch_output_format_xlsx_from_google_sheet(url: str) -> io.BytesIO:
    """表シュミの出力(入力データ／区間別ランナー／区間別配車の3シート)と同じ形式で
    作られたGoogleスプレッドシートを、そのままxlsxとして取得する。
    参加者の宿泊/運転/大/山/希望区間などはセルの値ではなく塗り色で符号化されている
    (_write_input_sheet参照)ため、色情報が失われるCSVエクスポートでは読み取れない。
    xlsx形式でのフルエクスポートなら塗り色ごと保持されるので、これを使う。
    「リンクを知っている全員が閲覧可」に共有されていれば認証なしで取得できる。"""
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", url)
    if not match:
        raise ValueError(f"URLからスプレッドシートIDを取得できませんでした: {url}")
    sheet_id = match.group(1)

    export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    resp = requests.get(export_url, timeout=30)
    content_type = resp.headers.get("content-type", "")
    if resp.status_code != 200 or "spreadsheet" not in content_type:
        raise ValueError(
            "スプレッドシートを認証なしで読み込めませんでした。"
            "共有設定が「リンクを知っている全員が閲覧可」になっているか確認してください。"
            f"(status={resp.status_code})"
        )
    return io.BytesIO(resp.content)


def _normalize_name(name: Optional[str]) -> str:
    """個人別まとめの数式(_normalize_expr)と同じ表記ゆれ吸収をPython側で行う。
    「山田太郎（4区も走る）」→「山田太郎」、「山田　太郎」→「山田太郎」。"""
    if not name:
        return ""
    name = re.split(r"[（(]", name)[0]
    return name.replace(" ", "").replace("　", "")


def _fill_rgb(cell) -> Optional[str]:
    fill = cell.fill
    if fill is None or fill.fgColor is None:
        return None
    return fill.fgColor.rgb


def import_participants_from_output_xlsx(path: str) -> Dict[str, Participant]:
    """入力データシートからParticipant一覧を復元する。"""
    wb = load_workbook(path)
    ws = wb["入力データ"]

    header = [ws.cell(row=1, column=c).value for c in range(1, len(INPUT_HEADER) + 1)]
    col = {name: i + 1 for i, name in enumerate(header) if name in INPUT_HEADER}

    section_cols = [col[str(i)] for i in range(1, 11)]

    participants: Dict[str, Participant] = {}
    p_index = 0
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=col["名前"]).value
        if not name:
            continue
        p_index += 1
        p_id = f"p{p_index}"

        grade_val = ws.cell(row=r, column=col["学年"]).value
        grade = int(grade_val) if grade_val is not None else 1

        stay_fill = _fill_rgb(ws.cell(row=r, column=col["宿泊"]))
        drive_fill = _fill_rgb(ws.cell(row=r, column=col["運転"]))
        large_fill = _fill_rgb(ws.cell(row=r, column=col["大"]))
        mountain_fill = _fill_rgb(ws.cell(row=r, column=col["山"]))

        preferred = [False] * 10
        priority = [False] * 10
        for i, c in enumerate(section_cols):
            fill = _fill_rgb(ws.cell(row=r, column=c))
            if fill == FILL_RED:
                preferred[i] = True
                priority[i] = True
            elif fill == FILL_BLACK:
                preferred[i] = True

        remaining_val = ws.cell(row=r, column=col["希望区間数"]).value
        remaining = int(remaining_val) if remaining_val is not None else sum(preferred)

        leave_val = ws.cell(row=r, column=col["離脱区間"]).value
        leaves_after = int(leave_val) if leave_val is not None else None

        participants[p_id] = Participant(
            id=p_id,
            name=str(name),
            preferred_sections=preferred,
            can_drive=(drive_fill == FILL_BLACK),
            can_drive_large=(large_fill == FILL_BLACK),
            can_drive_mountain=(mountain_fill == FILL_BLACK),
            staying_overnight=(stay_fill == FILL_BLACK),
            grade=grade,
            remaining_sections=remaining,
            leaves_after_section=leaves_after,
            priority_sections=priority,
        )
    return participants


def _name_to_id_map(participants: Dict[str, Participant]) -> Dict[str, str]:
    """正規化済み名前 -> participant_id。同姓同名がいる場合は最初の1人を使う
    (このプロジェクトの規模では実運用上ほぼ起こらない想定。起きた場合は
    呼び出し側で参加者一覧を見て名前を区別してもらう必要がある)。"""
    mapping: Dict[str, str] = {}
    for pid, p in participants.items():
        norm = _normalize_name(p.name)
        mapping.setdefault(norm, pid)
    return mapping


def _resolve_name(name_to_id: Dict[str, str], raw_name: Optional[str]) -> Optional[str]:
    if not raw_name:
        return None
    return name_to_id.get(_normalize_name(raw_name))


def import_plan_from_output_xlsx(path: str, participants: Dict[str, Participant]) -> List[SectionState]:
    """区間別配車・区間別ランナーシートからSectionStateのリストを復元する。
    表記ゆれのある名前や、手動編集で一部が空欄になった行にもある程度耐える
    (突き合わせられない名前は警告を出してスキップする)。"""
    wb = load_workbook(path)
    ws_runner = wb["区間別ランナー"]
    ws_cars = wb["区間別配車"]
    name_to_id = _name_to_id_map(participants)

    unmatched: List[str] = []

    def resolve(raw_name):
        pid = _resolve_name(name_to_id, raw_name)
        if pid is None and raw_name:
            unmatched.append(str(raw_name))
        return pid

    # --- 区間別ランナー: 1区間=1行 ---
    runners_by_label: Dict[str, List[str]] = {}
    runner_last_col = max(
        (c for c in range(2, ws_runner.max_column + 1)
         if str(ws_runner.cell(row=1, column=c).value or "").startswith("ランナー")),
        default=1,
    )
    for r in range(2, ws_runner.max_row + 1):
        label = ws_runner.cell(row=r, column=1).value
        if not label:
            continue
        ids = []
        for c in range(2, runner_last_col + 1):
            pid = resolve(ws_runner.cell(row=r, column=c).value)
            if pid:
                ids.append(pid)
        runners_by_label[label] = ids

    # --- 区間別配車: 区間ラベルはブロック先頭行にのみ入っている(結合セル) ---
    cars_header = [ws_cars.cell(row=1, column=c).value for c in range(1, ws_cars.max_column + 1)]
    driver_col = cars_header.index("運転手") + 1
    passenger_cols = [i + 1 for i, v in enumerate(cars_header) if v and str(v).startswith("同乗者")]
    mountain_col = cars_header.index("山行き") + 1 if "山行き" in cars_header else None

    title_row = None
    for r in range(2, ws_cars.max_row + 1):
        v = ws_cars.cell(row=r, column=1).value
        if isinstance(v, str) and v.startswith("■"):
            title_row = r
            break
    last_block_row = (title_row - 3) if title_row else ws_cars.max_row

    cars_by_label: Dict[str, List[CarState]] = {}
    current_label = None
    for r in range(2, last_block_row + 1):
        label_cell = ws_cars.cell(row=r, column=1).value
        if label_cell:
            current_label = label_cell
        if current_label is None:
            continue
        car_id = ws_cars.cell(row=r, column=2).value
        if not car_id:
            continue
        driver_pid = resolve(ws_cars.cell(row=r, column=driver_col).value)
        passenger_ids = [
            pid for c in passenger_cols
            if (pid := resolve(ws_cars.cell(row=r, column=c).value)) is not None
        ]
        # 「山行き」列は★山行き/🏨ホテル組/空欄のいずれかを表示するだけの列だが、
        # 元の表示をそのまま(変更せず)書き出し直せるように、その2値をここで読み戻す。
        mountain_text = str(ws_cars.cell(row=r, column=mountain_col).value or "").strip() if mountain_col else ""
        car_state = CarState(
            car_id=str(car_id),
            driver_id=driver_pid or "NO_DRIVER",
            passenger_ids=passenger_ids,
            car_type=CAR_TYPE.get(str(car_id), "normal"),
            is_mountain_goer=(mountain_text == "★山行き"),
            group=("hotel" if mountain_text == "🏨ホテル組" else None),
        )
        cars_by_label.setdefault(current_label, []).append(car_state)

    if unmatched:
        uniq = sorted(set(unmatched))
        print(f"  ⚠️ 参加者一覧に見つからなかった名前があります(スキップしました): {', '.join(uniq)}")

    plan: List[SectionState] = []
    for s in list(range(1, 11)) + [11]:
        label = section_label(s)
        plan.append(SectionState(
            section_id=s,
            runner_ids=runners_by_label.get(label, []),
            cars=cars_by_label.get(label, []),
        ))
    return plan
